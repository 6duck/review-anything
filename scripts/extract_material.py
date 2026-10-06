#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_material.py — unified extraction of study material (PDF / PPTX / DOCX / TXT / MD)

Usage:
    # Recon: which extraction route should this material take?
    python extract_material.py file.pdf --probe

    # Text-based PDF → plain text
    python extract_material.py file.pdf --out ./_extract/src.txt

    # Scanned PDF → render to PNG, hand to a vision model for transcription
    python extract_material.py scan.pdf --render-png ./_extract/pages/ --dpi 150
    python extract_material.py scan.pdf --render-png ./_extract/pages/ --pages 40-80

    # PPTX / DOCX
    python extract_material.py deck.pptx --out ./_extract/src.txt

Dependencies (install as needed; the script tells you what is missing):
    pip install pymupdf        # PDF
    pip install python-pptx    # PPTX
    pip install python-docx    # DOCX

Exit codes:
    0 = success. For --probe, 0 means "recon done"; it does NOT mean the material is fine.
    1 = failure (missing dependency / missing PNGs in the target page range / file does not exist / unsupported type)
    2 = partial (--render-png covered only part of the PDF; the remaining pages are not rendered yet)

Path note:
    Examples always use the relative path ./_extract/. /tmp/ exists only on Linux/macOS
    and fails on Windows — this script claims three-platform support, so do not hardcode /tmp/.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

SCAN_PAGE_CHAR_THRESHOLD = 50   # a page with fewer extractable chars than this is treated as an image/scanned page
NEAR_ZERO_CHARS = 5             # cutoff for a near-zero-char page
SCAN_RATIO_HIGH = 0.50          # share of near-zero-char pages > this → scanned (route B)
SCAN_RATIO_LOW = 0.10           # share of near-zero-char pages < this → text-based (route A)
BOOK_AVG_CHARS = 300            # average chars/page ≥ this → book-like; otherwise slide-like (affects wording only)
DEFAULT_BATCH = 3               # number of separate images handed to the vision model at once (stability first)


def die(msg: str, code: int = 1):
    print(msg, file=sys.stderr)
    sys.exit(code)


def parse_pages(spec: str, total: int):
    """'40-80' / '12' / '1,3,5-9' → list of page numbers (1-based)"""
    out = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            out.extend(range(int(a), int(b) + 1))
        else:
            out.append(int(part))
    return [p for p in out if 1 <= p <= total]


def require(mod: str, hint: str):
    try:
        return __import__(mod)
    except ImportError:
        die(f"missing dependency: {mod}\n  → {hint}")


def classify_route(chars) -> tuple[str, float]:
    """Single primary criterion: the share of near-zero-char pages.

    Why not use "average chars per page": a CJK textbook has 700–1000 chars per page,
    so "average > 100" always holds and has no discriminating power for books; a slide
    deck exported to PDF has only a few dozen chars per page, yet it does have a full
    text layer. Only "how many pages have no text layer at all" truly decides the route.
    """
    zero = sum(1 for c in chars if c < NEAR_ZERO_CHARS)
    ratio = zero / max(len(chars), 1)
    if ratio < SCAN_RATIO_LOW:
        return "A", ratio
    if ratio > SCAN_RATIO_HIGH:
        return "B", ratio
    return "C", ratio


# ---------------------------------------------------------------- PDF


def probe_pdf(path: Path):
    pymupdf = require("pymupdf", "pip install pymupdf")
    doc = pymupdf.open(str(path))
    n = len(doc)
    chars, scans, imgs = [], [], []
    for i in range(n):
        page = doc[i]
        t = page.get_text().strip()
        chars.append(len(t))
        if len(t) < SCAN_PAGE_CHAR_THRESHOLD:
            scans.append(i + 1)
        imgs.append(len(page.get_images()))

    avg = sum(chars) // max(n, 1)
    zero = sum(1 for c in chars if c < NEAR_ZERO_CHARS)
    route, ratio = classify_route(chars)
    kind = "book-like" if avg >= BOOK_AVG_CHARS else "slide-like"

    print(f"File        : {path.name}")
    print(f"Pages       : {n}")
    print(f"Near-zero   : {zero} pages ({100*zero//max(n,1)}%)  ← primary criterion")
    form = f"; kind: {kind}" if route == "A" else ""
    print(f"Avg chars/pg: {avg} chars (auxiliary info, not a criterion{form})")
    print(f"Scan-like   : {len(scans)} pages" + (f"  eg: {scans[:20]}" if scans else ""))
    print(f"Embedded img: first 5 pages {imgs[:5]}")

    if route == "A":
        print("\n→ Route A: text-based")
        print("   Extract directly; formulas in the text layer are often fragmented — check suspects against the original image:")
        print(f"   python {Path(__file__).name} {path.name} --out ./_extract/src.txt")
        if kind == "slide-like":
            print("   Note: few chars per page (slide-like), but still a text layer — route A is fine")
        print("\n→ Always spot-check after extraction: hyphenated line breaks / headers and footers / column order / formulas")
    elif route == "B":
        print("\n→ Route B: scanned / image-based (near-zero-char pages "
              f"{100*zero//max(n,1)}% > {int(SCAN_RATIO_HIGH*100)}%)")
        print("   Render PNGs and hand them to a vision model for transcription:")
        print(f"   python {Path(__file__).name} {path.name} "
              f"--render-png ./_extract/pages/ --dpi 150")
        print(f"   ⚠️ send {DEFAULT_BATCH} **separate** images at a time; never stitch two pages into one (stitched images get downsampled)")
        print("   ⚠️ for 200+ pages, shard and parallelize + flush every 20–30 pages and open a fresh window; do not run one worker to the end")
        print("   Prompt: transcribe these pages line by line and in full (output per page, separated by ===== p<N> =====).")
        print("           use LaTeX for formulas, keep numbering (definitions/theorems/examples) as-is, keep proofs complete.")
        print("   ⚠️ back-compute the page offset: PDF page = printed page + offset (do not trust PDF bookmarks)")
        print("      if a page has no printed page number (common on chapter openers) → infer from a neighboring numbered page ±1 and record the exception")
    else:
        print(f"\n→ Route C: mixed ({zero}/{n} pages without a text layer)")
        print(f"   body pages use A: python {Path(__file__).name} {path.name} --out ./_extract/src.txt")
        if scans:
            print(f"   no-text-layer pages use B: python {Path(__file__).name} {path.name} "
                  f"--render-png ./_extract/pages/ --dpi 150 --pages {scans[0]}-{scans[-1]}")
        print("\n→ Always spot-check after extraction: hyphenated line breaks / headers and footers / column order / formulas (formulas in the text layer are often fragmented)")

    return 0          # 0 = recon done; unrelated to whether the material is fine


def extract_pdf(path: Path, out: Path | None, pages_spec: str | None):
    pymupdf = require("pymupdf", "pip install pymupdf")
    doc = pymupdf.open(str(path))
    total = len(doc)
    targets = parse_pages(pages_spec, total) if pages_spec else list(range(1, total + 1))

    chunks, scan_pages, done = [], [], []
    for pno in targets:
        page = doc[pno - 1]
        text = page.get_text()
        if len(text.strip()) < SCAN_PAGE_CHAR_THRESHOLD:
            scan_pages.append(pno)
            text = "[[no extractable text layer on this page — likely scanned, needs vision transcription]]\n"
        chunks.append(f"\n\n===== p{pno} =====\n{text}")
        done.append(pno)

    body = "".join(chunks)
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(body, encoding="utf-8")
        print(f"Wrote: {out} ({len(body)} chars, {len(done)} pages)")
    else:
        print(body)

    print(f"Page reconciliation: target {len(targets)} pages · actually extracted {len(done)} pages"
          + (" (match ✓)" if len(done) == len(targets) else " (❌ mismatch)"))

    if scan_pages:
        print(f"\n⚠️ {len(scan_pages)} pages have no text layer and need the vision route: {scan_pages[:30]}"
              f"{' …' if len(scan_pages) > 30 else ''}", file=sys.stderr)
        print("   → render these pages with --render-png, then hand them to a vision model", file=sys.stderr)

    return 0 if len(done) == len(targets) else 1


def render_pdf(path: Path, outdir: Path, dpi: int, pages_spec: str | None,
               batch: int = DEFAULT_BATCH):
    pymupdf = require("pymupdf", "pip install pymupdf")
    doc = pymupdf.open(str(path))
    total = len(doc)
    targets = parse_pages(pages_spec, total) if pages_spec else list(range(1, total + 1))
    if not targets:
        die(f"page range is empty or entirely out of bounds (PDF has {total} pages)")
    outdir.mkdir(parents=True, exist_ok=True)

    rendered, failed, size = [], [], None
    for pno in targets:
        try:
            pix = doc[pno - 1].get_pixmap(dpi=dpi)
            size = f"{pix.width}x{pix.height}px"
            pix.save(str(outdir / f"p{pno:04d}.png"))
            rendered.append(pno)
        except Exception as e:                       # a single-page failure must not sink the whole batch
            failed.append((pno, str(e)))

    existing = sorted(outdir.glob("p[0-9][0-9][0-9][0-9].png"))
    existing_pages = {int(p.stem[1:]) for p in existing}

    # ---- page reconciliation: the only way to prevent "thought rendering was done, actually 100 images short"
    print(f"\nPage reconciliation: PDF has {total} pages · this run targets {len(targets)} pages "
          f"· {len(rendered)} rendered successfully · {len(existing)} PNGs currently in the dir")
    print(f"          output dir: {outdir} (dpi={dpi}, {size})")
    if failed:
        for pno, err in failed[:10]:
            print(f"  ❌ p{pno} render failed: {err}", file=sys.stderr)

    missing_in_range = [p for p in targets if p not in existing_pages]
    if missing_in_range:
        print(f"❌ {len(missing_in_range)} pages in the target range still have no PNG: {missing_in_range[:20]}",
              file=sys.stderr)
        print("   Do not start transcription — missing pages cause silently missing chapters in the notes.", file=sys.stderr)
        return 1

    # full-coverage check: piecewise rendering is legitimate, but silently being "100 images short" is not
    uncovered = [p for p in range(1, total + 1) if p not in existing_pages]
    if uncovered:
        print(f"⚠️ the dir covers only {total - len(uncovered)}/{total} pages of the PDF; "
              f"{len(uncovered)} pages are not rendered yet (eg: {uncovered[:20]}"
              f"{' …' if len(uncovered) > 20 else ''})")
        print("   → this is fine to run for piecewise rendering, but **confirm every range is fully rendered before transcription**, "
              "otherwise the notes will silently miss chapters.", file=sys.stderr)
        partial = True
    else:
        print(f"✓ Full coverage: all {total} PDF pages have PNGs")
        partial = False

    # ---- batch manifest: hand batch **separate** images to the vision model at once
    groups = [existing[i:i + batch] for i in range(0, len(existing), batch)]
    print(f"\nHand {batch} separate images per batch to the vision model ({len(groups)} batches total):")
    for i, g in enumerate(groups[:5], 1):
        print(f"  batch {i:>3}: {' '.join(p.name for p in g)}")
    if len(groups) > 5:
        print(f"  …({len(groups) - 5} more batches, see the manifest)")

    manifest = outdir / "batches.txt"
    manifest.write_text("\n".join(" ".join(str(p) for p in g) for g in groups), encoding="utf-8")
    print(f"\nBatch manifest: {manifest} (one batch per line, convenient for batch scheduling)")
    print(f"⚠️ {batch} **separate** images at a time ≠ stitching two pages into one: stitched images get downsampled, blurring out subscripts.")
    print("⚠️ for 200+ pages, shard and parallelize (split into 2–4 non-overlapping ranges, named pXXX-YYY.md), "
          "flushing every 20–30 pages and opening a fresh window.")
    return 2 if partial else 0


# ---------------------------------------------------------------- PPTX


def extract_pptx(path: Path, out: Path | None):
    require("pptx", "pip install python-pptx")
    from pptx import Presentation

    prs = Presentation(str(path))
    chunks = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        lines.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[Notes] {notes}")
        chunks.append(f"\n\n===== slide{i} =====\n" + "\n".join(lines))

    body = "".join(chunks)
    n_slides = len(chunks)
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(body, encoding="utf-8")
        print(f"Wrote: {out} ({n_slides} slides, {len(body)} chars)")
    else:
        print(body)
    print("⚠️ extraction follows shape z-order, which is not the visual reading order → you must reorder after extraction", file=sys.stderr)
    print("⚠️ charts/diagrams cannot be extracted → when there are many, export page images and use the vision route", file=sys.stderr)
    return 0


# ---------------------------------------------------------------- DOCX


def extract_docx(path: Path, out: Path | None):
    require("docx", "pip install python-docx")
    import docx

    d = docx.Document(str(path))
    lines = []
    for para in d.paragraphs:
        if para.text.strip():
            style = para.style.name if para.style else ""
            prefix = "# " if style.startswith("Heading 1") else ""
            lines.append(prefix + para.text)
    for ti, table in enumerate(d.tables, start=1):
        lines.append(f"\n[table {ti}]")
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    body = "\n".join(lines)
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(body, encoding="utf-8")
        print(f"Wrote: {out} ({len(body)} chars)")
    else:
        print(body)
    return 0


# ---------------------------------------------------------------- main


def main(argv=None):
    ap = argparse.ArgumentParser(description="Unified extraction of study material")
    ap.add_argument("input", help="material file")
    ap.add_argument("--out", help="output text file (prints to stdout if omitted)")
    ap.add_argument("--probe", action="store_true", help="recon only: decide which extraction route to take")
    ap.add_argument("--render-png", metavar="DIR", help="render PDF pages to PNG in this directory")
    ap.add_argument("--dpi", type=int, default=150, help="render resolution (default 150, the sweet spot)")
    ap.add_argument("--pages", help="page range, e.g. 40-80 or 1,3,5-9 (all pages by default)")
    ap.add_argument("--batch", type=int, default=DEFAULT_BATCH,
                    help=f"separate images per batch handed to the vision model (default {DEFAULT_BATCH}, stability first)")
    args = ap.parse_args(argv)

    path = Path(args.input)
    if not path.is_file():
        die(f"file does not exist: {path}")

    suffix = path.suffix.lower()

    if args.probe:
        if suffix == ".pdf":
            return probe_pdf(path)
        if suffix in (".pptx", ".ppt"):
            return extract_pptx(path, None)
        if suffix in (".docx", ".doc"):
            return extract_docx(path, None)
        print(f"{path.name}: text file, read it directly")
        return 0

    out = Path(args.out) if args.out else None

    if suffix == ".pdf":
        if args.render_png:
            return render_pdf(path, Path(args.render_png), args.dpi, args.pages, args.batch)
        return extract_pdf(path, out, args.pages)
    if suffix in (".pptx", ".ppt"):
        return extract_pptx(path, out)
    if suffix in (".docx", ".doc"):
        return extract_docx(path, out)
    if suffix in (".txt", ".md", ".csv", ".json"):
        body = path.read_text(encoding="utf-8", errors="replace")
        if out:
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_text(body, encoding="utf-8")
            print(f"Wrote: {out} ({len(body)} chars)")
        else:
            print(body)
        return 0

    die(f"unsupported type: {suffix}\nSupported: .pdf .pptx .docx .txt .md .csv .json")


if __name__ == "__main__":
    sys.exit(main())
